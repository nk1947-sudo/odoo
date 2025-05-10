from odoo.tests.common import TransactionCase
from datetime import date
from odoo.exceptions import UserError, ValidationError
from openpyxl import load_workbook
from io import BytesIO


class TestWorkReport(TransactionCase):

    def setUp(self):
        super().setUp()
        self.employee = self.env['hr.employee'].create({'name': 'Test Employee'})
        self.project = self.env['project.project'].create({'name': 'Test Project'})
        self.task = self.env['project.task'].create({
            'name': 'Test Task',
            'project_id': self.project.id,
        })
        self.template = self.env['work.report.template'].create({
            'name': 'Weekly Template',
            'line_ids': [(0, 0, {
                'name': 'Dev Task',
                'project_id': self.project.id,
                'task_id': self.task.id,
                'billable': True
            })]
        })

    def test_create_report_with_template(self):
        report = self.env['work.report'].create({
            'date': date.today(),
            'employee_id': self.employee.id,
            'template_id': self.template.id
        })
        self.assertTrue(report.line_ids)
        self.assertEqual(report.state, 'draft')

    def test_approval_flow(self):
        report = self.env['work.report'].create({
            'date': date.today(),
            'employee_id': self.employee.id,
            'line_ids': [(0, 0, {
                'name': 'Work',
                'project_id': self.project.id,
                'task_id': self.task.id,
                'hours': 8,
                'billable': True
            })]
        })
        report.action_submit()
        self.assertEqual(report.state, 'submitted')

        report.action_manager_approve()
        self.assertEqual(report.state, 'manager_approval')

        report.action_approve()
        self.assertEqual(report.state, 'approved')
        self.assertEqual(report.approved_by, self.env.user)

    def test_compute_hours_and_efficiency(self):
        report = self.env['work.report'].create({
            'date': date.today(),
            'employee_id': self.employee.id,
            'line_ids': [
                (0, 0, {'name': 'Billable', 'hours': 6, 'billable': True}),
                (0, 0, {'name': 'Non-Billable', 'hours': 2, 'billable': False})
            ]
        })
        self.assertEqual(report.total_hours, 8)
        self.assertEqual(report.billable_hours, 6)
        self.assertEqual(report.non_billable_hours, 2)
        self.assertEqual(report.efficiency, 75)

    def test_validation_errors(self):
        report = self.env['work.report'].create({
            'date': date.today(),
            'employee_id': self.employee.id,
        })
        with self.assertRaises(UserError):
            report.action_submit()

        with self.assertRaises(ValidationError):
            self.env['work.report.line'].create({
                'report_id': report.id,
                'name': 'Invalid',
                'hours': -5
            })

    def test_timesheet_creation(self):
        report = self.env['work.report'].create({
            'date': date.today(),
            'employee_id': self.employee.id,
            'line_ids': [(0, 0, {
                'name': 'Work',
                'project_id': self.project.id,
                'task_id': self.task.id,
                'hours': 4,
                'billable': True
            })]
        })
        line = report.line_ids[0]
        line.create_timesheet()

        timesheet = self.env['account.analytic.line'].search([
            ('project_id', '=', self.project.id),
            ('task_id', '=', self.task.id),
            ('unit_amount', '=', 4)
        ])
        self.assertTrue(timesheet)

    def test_export_xlsx_report_content(self):
        """Optional XLSX export test if _generate_xlsx_report() is defined."""
        report = self.env['work.report'].create({
            'date': date.today(),
            'employee_id': self.employee.id,
            'line_ids': [(0, 0, {
                'name': 'XLSX Test Line',
                'project_id': self.project.id,
                'task_id': self.task.id,
                'hours': 2,
                'billable': True
            })]
        })

        if hasattr(report, '_generate_xlsx_report'):
            binary_data = report._generate_xlsx_report()
            wb = load_workbook(filename=BytesIO(binary_data), data_only=True)
            ws = wb.active

            # Ensure test line is present in report
            found = any('XLSX Test Line' in str(cell) for row in ws.iter_rows(values_only=True) for cell in row if cell)
            self.assertTrue(found, "Expected line not found in XLSX export")
